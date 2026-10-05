"""ترجمة المستند جدولَ Excel للمراجعة: جملة في كل صف بنوعها ونصها الأصلي وترجمتها وحالتها."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from classify.services.moneer_ui import attribution_label, title_label

HEADERS = ("#", "النوع", "النص الأصلي", "الترجمة", "الحالة")
WIDTHS = {"A": 6, "B": 12, "C": 60, "D": 60, "E": 16}

# تسميات الأنواع كما تعرضها الواجهة
_KIND_LABELS = {
    "quran": "آية",
    "hadith": "حديث",
    "athar": "أثر",
    "term": "مصطلح",
    "citation": "عزو",
}
_TYPE_LABELS = {"ayah": "آية", "hadith": "حديث", "athar": "أثر"}


def kind_label(phrase) -> str:
    analysis = getattr(phrase, "analysis", None)
    kind = analysis.kind if analysis else ""
    if kind == "heading" or phrase.tag.startswith("heading"):
        return title_label(phrase.text)
    if kind == "attribution":
        return attribution_label(phrase.text)
    if phrase.tag == "footnote":
        return "هامش"
    return _KIND_LABELS.get(kind) or _TYPE_LABELS.get(phrase.content_type.code, "نص عام")


def status_label(phrase, row) -> str:
    if not phrase.translatable:
        return "يُنقل كما هو"
    if row is None or not row.translation.strip():
        return "بلا ترجمة"
    return row.get_status_display()


def build_xlsx(document_translation) -> tuple[bytes, int, int]:
    """``(المحتوى، عدد الجمل، القابلة للترجمة بلا ترجمة)``؛ كل جمل المستند بترتيبه ولو بلا ترجمة."""
    translations = {row.phrase_id: row for row in document_translation.phrases.all()}
    phrases = document_translation.document.phrases.select_related(
        "content_type", "analysis"
    ).order_by("group_id", "group_order")
    ltr = document_translation.target_language.direction == "ltr"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "الترجمة"
    sheet.sheet_view.rightToLeft = True
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2464EC")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column, width in WIDTHS.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"

    count = missing = 0
    for count, phrase in enumerate(phrases, 1):
        row = translations.get(phrase.pk)
        translation = row.translation.strip() if row else ""
        if phrase.translatable and not translation:
            missing += 1
        sheet.append(
            [count, kind_label(phrase), phrase.text, translation, status_label(phrase, row)]
        )
        line = sheet[sheet.max_row]
        for cell in line:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        line[3].alignment = Alignment(
            vertical="top", wrap_text=True, horizontal="left" if ltr else "right"
        )

    out = BytesIO()
    workbook.save(out)
    return out.getvalue(), count, missing
