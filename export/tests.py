"""اختبارات التصدير: الافتراضيات وتحويل الصفوف وتوليد DOCX وPDF والتنزيل الآني والأوامر وadmin."""

import shutil
import tempfile
import zipfile
from io import StringIO
from pathlib import Path

from django.contrib import admin
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.core.management.base import CommandError
from django.db import IntegrityError, transaction
from django.test import TestCase, override_settings
from django.urls import reverse
from spire.doc import Document as SpireDocument

from content.models import Document, DocumentTranslation, Phrase, PhraseTranslation
from core.models import ContentTypeLookup, Language
from users.models import User

from .defaults import LTR_DIR, OPTION_KEYS, RTL_DIR, defaults_for, validate_options
from .models import ExportFormat
from .services.export import CONTENT_TYPES, ExportError, export_translation
from .services.rows import build_rows

OUT = tempfile.mkdtemp(prefix="moneer-export-tests-")


def tearDownModule():
    shutil.rmtree(OUT, ignore_errors=True)


def write(content, name):
    path = Path(OUT) / name
    path.write_bytes(content)
    return path


def read_docx(path):
    """(قائمة (نمط، نص) لكل فقرة غير فارغة في كل الأقسام، نص رأس كل قسم مجموعاً)."""
    document = SpireDocument()
    document.LoadFromFile(str(path))
    paragraphs, headers = [], []
    for i in range(document.Sections.Count):
        section = document.Sections[i]
        header = section.HeadersFooters.Header
        texts = [header.Paragraphs[k].Text.strip() for k in range(header.Paragraphs.Count)]
        headers.append(" | ".join(text for text in texts if text))
        for j in range(section.Paragraphs.Count):
            paragraph = section.Paragraphs[j]
            if paragraph.Text.strip():
                paragraphs.append((paragraph.StyleName, paragraph.Text.strip()))
    document.Close()
    return paragraphs, headers


@override_settings(GEMINI_API_KEY="", OPENAI_API_KEY="")
class ExportTestCase(TestCase):
    """أساس مشترك: لغتان وأنواع محتوى ومستند بجمل تغطي كل مسارات المولّد وترجمته الإنجليزية."""

    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user(
            email="exporter@example.com", password="x", full_name="مصدّر تجريبي"
        )
        cls.english, _ = Language.objects.get_or_create(
            iso_code="en",
            defaults={"name": "الإنجليزية", "name_en": "English", "direction": "ltr"},
        )
        cls.urdu, _ = Language.objects.get_or_create(
            iso_code="ur",
            defaults={"name": "الأردية", "name_en": "Urdu", "direction": "rtl"},
        )
        cls.text_type = ContentTypeLookup.objects.get(code="text")
        cls.ayah_type = ContentTypeLookup.objects.get(code="ayah")
        cls.hadith_type = ContentTypeLookup.objects.filter(code="hadith").first() or cls.text_type
        cls.document = Document.objects.create(
            title="خطبة تجريبية: التصدير", kind=Document.Kind.KHUTBAH
        )
        cls.phrases = {}
        rows = [
            # (مفتاح، فقرة، ترتيب، نص، وسم، نوع نص، نوع محتوى، قابل للترجمة، الترجمة)
            ("title", 1, 1, "عنوان الخطبة", "heading1", "h", cls.text_type, True, "Khutbah Title"),
            (
                "p1",
                2,
                1,
                "الحمد لله رب العالمين[1]",
                "p",
                "p",
                cls.text_type,
                True,
                "Praise be to Allah, Lord of the worlds.[1]",
            ),
            ("p2", 2, 2, "نحمده ونستعينه", "p", "p", cls.text_type, True, "We praise Him."),
            (
                "note",
                3,
                1,
                "[1] نص الهامش",
                "footnote",
                "footnote",
                cls.text_type,
                True,
                "[1] Footnote text here.",
            ),
            (
                "aya",
                4,
                1,
                "﴿قُلْ هُوَ اللَّهُ أَحَدٌ﴾",
                "p",
                "p",
                cls.ayah_type,
                True,
                "Say: He is Allah, the One (1)",
            ),
            ("ref", 4, 2, "[الإخلاص: 1]", "p", "p", cls.text_type, False, None),
            ("h2", 5, 1, "الخطبة الثانية", "heading2", "h", cls.text_type, True, "Second Sermon"),
            ("naqhara", 6, 1, "سبحان الله", "p", "naqhara", cls.text_type, True, "Subhanallah"),
            ("missing", 7, 1, "جملة بلا ترجمة", "p", "p", cls.text_type, True, None),
            (
                "hadith",
                8,
                1,
                "«إنما الأعمال بالنيات»",
                "p",
                "p",
                cls.hadith_type,
                True,
                "Deeds are by intentions.",
            ),
        ]
        cls.dt = DocumentTranslation.objects.create(
            document=cls.document,
            target_language=cls.english,
            status=DocumentTranslation.Status.REVIEW,
            created_by=cls.user,
        )
        for key, group, order, text, tag, text_type, content_type, translatable, trans in rows:
            phrase = Phrase.objects.create(
                document=cls.document,
                content_type=content_type,
                group_id=group,
                group_order=order,
                text=text,
                tag=tag,
                text_type=text_type,
                translatable=translatable,
            )
            cls.phrases[key] = phrase
            if trans:
                PhraseTranslation.objects.create(
                    document_translation=cls.dt,
                    phrase=phrase,
                    ai_translation=trans,
                    translation=trans,
                    status=PhraseTranslation.Status.SUGGESTED,
                )
        cls.english_format = ExportFormat.objects.create(language=cls.english, name="افتراضي")


class DefaultsTests(TestCase):
    def test_direction_aware_defaults_share_the_same_keys(self):
        rtl, ltr = defaults_for("rtl"), defaults_for("ltr")
        self.assertEqual(set(rtl), set(ltr))
        self.assertEqual(set(rtl), OPTION_KEYS)
        self.assertEqual(rtl["main_text_dir"], RTL_DIR)
        self.assertEqual(ltr["main_text_dir"], LTR_DIR)
        self.assertEqual(rtl["heading_1_alignment"], "اليمين")
        self.assertEqual(ltr["heading_1_alignment"], "اليسار")
        self.assertEqual(rtl["paragraph_fonts"], "Traditional Arabic")
        self.assertEqual(ltr["paragraph_fonts"], "Arial")
        # النص العربي بخطوط KFGQPC في الاتجاهين
        self.assertTrue(ltr["aya_fonts"].startswith("KFGQPC"))
        self.assertTrue(rtl["arabic_title_fonts"].startswith("KFGQPC"))

    def test_validate_options(self):
        self.assertEqual(validate_options({}), [])
        self.assertEqual(validate_options({"paragraph_fonts": "Amiri", "include_index": False}), [])
        errors = validate_options({"bogus": 1, "include_index": "yes"})
        self.assertEqual(len(errors), 2)
        self.assertIn("bogus", errors[0])
        self.assertIn("include_index", errors[1])
        self.assertEqual(len(validate_options([])), 1)


class ExportFormatTests(ExportTestCase):
    def test_first_format_is_default_and_only_one_default_per_language(self):
        self.assertTrue(self.english_format.is_default)
        second = ExportFormat.objects.create(language=self.english, name="مطبوع", is_default=True)
        self.english_format.refresh_from_db()
        self.assertFalse(self.english_format.is_default)
        self.assertTrue(second.is_default)
        third = ExportFormat.objects.create(language=self.english, name="ثالث")
        self.assertFalse(third.is_default)
        self.assertEqual(
            ExportFormat.objects.filter(language=self.english, is_default=True).count(), 1
        )
        # لغة أخرى لها افتراضيها المستقل
        urdu = ExportFormat.objects.create(language=self.urdu, name="افتراضي")
        self.assertTrue(urdu.is_default)

    def test_database_constraint_rejects_a_second_default(self):
        second = ExportFormat.objects.create(language=self.english, name="مطبوع")
        with self.assertRaises(IntegrityError), transaction.atomic():
            ExportFormat.objects.filter(pk=second.pk).update(is_default=True)

    def test_name_is_unique_per_language(self):
        with self.assertRaises(IntegrityError), transaction.atomic():
            ExportFormat.objects.create(language=self.english, name=" افتراضي ")

    def test_preferences_merge_options_over_direction_defaults(self):
        fmt = ExportFormat.objects.create(
            language=self.urdu, name="أردو", options={"paragraph_fonts": "Jameel Noori Nastaleeq"}
        )
        prefs = fmt.preferences()
        self.assertEqual(prefs["paragraph_fonts"], "Jameel Noori Nastaleeq")
        self.assertEqual(prefs["main_text_dir"], RTL_DIR)
        self.assertEqual(set(prefs), OPTION_KEYS)

    def test_clean_rejects_unknown_option_keys(self):
        fmt = ExportFormat(language=self.english, name="خاطئ", options={"bogus": 1})
        with self.assertRaises(ValidationError) as ctx:
            fmt.full_clean()
        self.assertIn("options", ctx.exception.message_dict)


class RowsTests(ExportTestCase):
    def test_rows_follow_document_order_and_map_phrase_fields(self):
        rows, missing = build_rows(self.dt)
        self.assertEqual(missing, [self.phrases["missing"].pk])
        self.assertEqual(len(rows), 9)
        by_id = {row["id"]: row for row in rows}
        title = by_id[self.phrases["title"].pk]
        self.assertEqual((title["type"], title["tag"], title["splitted"]), ("title", "h1", 0))
        self.assertEqual(by_id[self.phrases["h2"].pk]["tag"], "h2")
        self.assertEqual(by_id[self.phrases["note"].pk]["type"], "footnote")
        self.assertEqual(by_id[self.phrases["naqhara"].pk]["type"], "naqhara")
        self.assertEqual(by_id[self.phrases["aya"].pk]["sub_type"], "aya")
        if self.hadith_type.code == "hadith":
            self.assertEqual(by_id[self.phrases["hadith"].pk]["sub_type"], "hadith")
        p1, p2 = by_id[self.phrases["p1"].pk], by_id[self.phrases["p2"].pk]
        self.assertEqual((p1["splitted"], p2["splitted"]), (0, 1))
        self.assertEqual(p1["split_group"], p2["split_group"])
        self.assertEqual(p2["translation"], "We praise Him.")
        self.assertEqual(p2["original_text"], "نحمده ونستعينه")
        ref = by_id[self.phrases["ref"].pk]
        self.assertEqual(ref["translation"], "[الإخلاص: 1]")  # غير القابل للترجمة بنصه العربي
        self.assertEqual(ref["splitted"], 1)
        self.assertEqual(
            [row["id"] for row in rows][:2], [self.phrases["title"].pk, self.phrases["p1"].pk]
        )

    def test_untranslatable_rows_can_be_excluded(self):
        rows, _ = build_rows(self.dt, include_untranslatable=False)
        self.assertNotIn(self.phrases["ref"].pk, [row["id"] for row in rows])
        self.assertEqual(len(rows), 8)


class DocxExportTests(ExportTestCase):
    def test_export_builds_a_real_docx_with_headings_footnote_and_index(self):
        self.english_format.options = {"add_header_title": True}
        self.english_format.save()

        result = export_translation(self.dt)

        self.assertEqual(result.kind, "docx")
        self.assertEqual((result.rows, result.missing), (9, 1))
        self.assertEqual(result.stats, {"footnotes_placed": 1, "footnotes_missed": 0})
        self.assertTrue(result.file_name.endswith(f"[en{self.dt.pk}].docx"), result.file_name)
        self.assertNotIn(":", result.file_name)  # الأحرف الممنوعة في أسماء الملفات تُستبدل
        self.assertEqual(result.content[:2], b"PK")  # DOCX ملف zip

        path = write(result.content, "en.docx")
        paragraphs, headers = read_docx(path)
        styles = {text: style for style, text in paragraphs}
        self.assertEqual(styles["Khutbah Title"], "Heading1")
        self.assertEqual(styles["Second Sermon"], "Heading2")
        body = "\n".join(text for _, text in paragraphs)
        self.assertIn("Praise be to Allah, Lord of the worlds. We praise Him.", body)
        self.assertNotIn("[1]", body)  # علامات الحواشي حُذفت بعد تعليقها
        self.assertIn("(قُلْ هُوَ اللَّهُ أَحَدٌ) Say: He is Allah, the One (1) [الإخلاص: 1]", body)
        self.assertIn("Deeds are by intentions.", body)
        self.assertNotIn("Subhanallah", body)  # النقحرة معطَّلة افتراضياً
        self.assertNotIn("جملة بلا ترجمة", body)  # الجملة بلا ترجمة أُسقطت
        self.assertIn("خطبة تجريبية: التصدير", body)  # العنوان العربي في صفحة العنوان
        self.assertIn("الفهرس", body)
        self.assertTrue(any(style == "IndexStyle" for style, _ in paragraphs))
        self.assertIn(f"en{self.dt.pk} - ", body)  # وسم الكتاب
        self.assertEqual(headers[1], "Khutbah Title")  # عنوان الكتاب في رأس قسم المحتوى
        self.assertIn(
            "Footnote text here.", zipfile.ZipFile(path).read("word/footnotes.xml").decode()
        )

    def test_export_pdf(self):
        result = export_translation(self.dt, kind="pdf")
        self.assertEqual(result.kind, "pdf")
        self.assertTrue(result.file_name.endswith(".pdf"))
        self.assertEqual(result.content[:5], b"%PDF-")

    def test_rtl_language_uses_rtl_defaults_and_builds(self):
        dt = DocumentTranslation.objects.create(
            document=self.document,
            target_language=self.urdu,
            status=DocumentTranslation.Status.REVIEW,
        )
        for key in ("title", "p1", "p2", "note", "aya", "hadith", "naqhara", "h2"):
            PhraseTranslation.objects.create(
                document_translation=dt,
                phrase=self.phrases[key],
                translation=f"اردو {key}[1]" if key == "p1" else f"اردو {key}",
                status=PhraseTranslation.Status.SUGGESTED,
            )
        fmt = ExportFormat.objects.create(
            language=self.urdu,
            name="افتراضي",
            options={
                "include_naqhara": True,
                "include_arabic_naqhara": True,
                "split_arabic_aya": True,
            },
        )
        self.assertEqual(fmt.preferences()["main_text_dir"], RTL_DIR)
        result = export_translation(dt)
        self.assertEqual(result.stats["footnotes_placed"], 1)
        body = "\n".join(text for _, text in read_docx(write(result.content, "ur.docx"))[0])
        self.assertIn("اردو title", body)
        self.assertIn("«سبحان الله»", body)  # النقحرة العربية بين قوسي اقتباس
        self.assertIn("﴿قُلْ هُوَ اللَّهُ أَحَدٌ﴾", body)  # الأقواس القرآنية تبقى في RTL

    def test_missing_default_format_raises(self):
        french, _ = Language.objects.get_or_create(
            iso_code="fr", defaults={"name": "الفرنسية", "name_en": "French", "direction": "ltr"}
        )
        dt = DocumentTranslation.objects.create(
            document=self.document, target_language=french, status=DocumentTranslation.Status.REVIEW
        )
        ExportFormat.objects.filter(language=french).delete()
        with self.assertRaises(ExportError) as ctx:
            export_translation(dt)
        self.assertIn("ensure_default_formats", str(ctx.exception))

    def test_unknown_kind_raises(self):
        with self.assertRaises(ExportError):
            export_translation(self.dt, kind="odt")


class XlsxExportTests(ExportTestCase):
    @staticmethod
    def read(content):
        from io import BytesIO

        from openpyxl import load_workbook

        return list(load_workbook(BytesIO(content)).active.iter_rows(values_only=True))

    def test_xlsx_lists_every_phrase_with_its_type_translation_and_status(self):
        result = export_translation(self.dt, kind="xlsx")
        rows = self.read(result.content)
        self.assertEqual(rows[0], ("#", "النوع", "النص الأصلي", "الترجمة", "الحالة"))
        self.assertEqual(len(rows) - 1, len(self.phrases))
        by_text = {row[2]: row for row in rows[1:]}
        self.assertEqual(
            by_text["عنوان الخطبة"][1:], ("عنوان", "عنوان الخطبة", "Khutbah Title", "مقترحة آلياً")
        )
        self.assertEqual(by_text["﴿قُلْ هُوَ اللَّهُ أَحَدٌ﴾"][1], "آية")
        self.assertEqual(by_text["[1] نص الهامش"][1], "هامش")
        self.assertFalse(by_text["[الإخلاص: 1]"][3])
        self.assertEqual(by_text["[الإخلاص: 1]"][4], "يُنقل كما هو")
        self.assertEqual(by_text["جملة بلا ترجمة"][4], "بلا ترجمة")
        self.assertEqual((result.rows, result.missing), (len(self.phrases), 1))
        self.assertTrue(result.file_name.endswith(".xlsx"))

    def test_xlsx_needs_no_export_format(self):
        ExportFormat.objects.all().delete()
        self.assertEqual(export_translation(self.dt, kind="xlsx").kind, "xlsx")


class DownloadViewTests(ExportTestCase):
    def url(self, kind="docx", pk=None):
        return reverse("export:download", args=[pk or self.dt.pk, kind])

    def test_docx_download_is_generated_on_the_fly(self):
        self.assertEqual(self.url(), f"/export/{self.dt.pk}/docx/")
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], CONTENT_TYPES["docx"])
        disposition = response["Content-Disposition"]
        self.assertTrue(disposition.startswith("attachment;"), disposition)
        self.assertIn(f"en{self.dt.pk}", disposition)
        self.assertEqual(b"".join(response.streaming_content)[:2], b"PK")

    def test_pdf_download(self):
        response = self.client.get(self.url("pdf"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], CONTENT_TYPES["pdf"])
        self.assertEqual(b"".join(response.streaming_content)[:5], b"%PDF-")

    def test_xlsx_download(self):
        response = self.client.get(self.url("xlsx"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], CONTENT_TYPES["xlsx"])
        self.assertIn(".xlsx", response["Content-Disposition"])
        self.assertEqual(b"".join(response.streaming_content)[:2], b"PK")

    def test_specific_format_must_belong_to_the_translation_language(self):
        other = ExportFormat.objects.create(
            language=self.english, name="بلا فهرس", options={"include_index": False}
        )
        urdu = ExportFormat.objects.create(language=self.urdu, name="أردو")
        self.assertEqual(self.client.get(self.url() + f"?format={other.pk}").status_code, 200)
        self.assertEqual(self.client.get(self.url() + f"?format={urdu.pk}").status_code, 404)
        self.assertEqual(self.client.get(self.url() + "?format=abc").status_code, 404)

    def test_unknown_translation_or_kind_is_404(self):
        self.assertEqual(self.client.get(self.url(pk=999999)).status_code, 404)
        self.assertEqual(self.client.get(self.url("odt")).status_code, 404)

    def test_missing_default_format_is_400_with_a_message(self):
        ExportFormat.objects.filter(language=self.english).delete()
        response = self.client.get(self.url())
        self.assertEqual(response.status_code, 400)
        self.assertIn("ensure_default_formats", response.content.decode())


class CommandTests(ExportTestCase):
    def test_ensure_default_formats_is_idempotent_and_skips_languages_with_formats(self):
        out = StringIO()
        call_command("ensure_default_formats", stdout=out)
        self.assertEqual(ExportFormat.objects.filter(language=self.english).count(), 1)
        self.assertTrue(ExportFormat.objects.filter(language=self.urdu, is_default=True).exists())
        self.assertEqual(
            ExportFormat.objects.filter(is_default=True).count(), Language.objects.count()
        )
        again = StringIO()
        call_command("ensure_default_formats", stdout=again)
        self.assertIn("أُنشئ 0", again.getvalue())

    def test_ensure_default_formats_for_selected_languages(self):
        call_command("ensure_default_formats", "--language", "UR", stdout=StringIO())
        self.assertTrue(ExportFormat.objects.filter(language=self.urdu).exists())
        self.assertEqual(
            ExportFormat.objects.exclude(language__in=[self.english, self.urdu]).count(), 0
        )

    def test_export_translation_command_writes_the_file(self):
        out = StringIO()
        call_command("export_translation", str(self.dt.pk), "--output", OUT, stdout=out)
        self.assertIn("صُدِّر 9 جملة", out.getvalue())
        generated = [
            p for p in Path(OUT).glob("*.docx") if p.name.endswith(f"[en{self.dt.pk}].docx")
        ]
        self.assertEqual(len(generated), 1)
        self.assertEqual(generated[0].read_bytes()[:2], b"PK")

        target = Path(OUT) / "book.pdf"
        call_command(
            "export_translation",
            str(self.dt.pk),
            "--kind",
            "pdf",
            "--output",
            str(target),
            stdout=out,
        )
        self.assertEqual(target.read_bytes()[:5], b"%PDF-")

        with self.assertRaises(CommandError):
            call_command("export_translation", "999999", stdout=StringIO())
        with self.assertRaises(CommandError):
            call_command("export_translation", str(self.dt.pk), "--format", "999999")


class AdminTests(ExportTestCase):
    def setUp(self):
        self.admin_user = User.objects.create_superuser(
            email="admin@example.com", password="x", full_name="مشرف"
        )
        self.client.force_login(self.admin_user)

    def test_export_format_is_registered_and_changelist_renders(self):
        self.assertIn(ExportFormat, admin.site._registry)
        response = self.client.get(reverse("admin:export_exportformat_changelist"))
        self.assertEqual(response.status_code, 200)

    def test_document_translation_admin_links_to_on_the_fly_export(self):
        response = self.client.get(reverse("admin:content_documenttranslation_changelist"))
        self.assertContains(response, f"/export/{self.dt.pk}/docx/")
        self.assertContains(response, f"/export/{self.dt.pk}/pdf/")
        self.assertContains(response, f"/export/{self.dt.pk}/xlsx/")
        response = self.client.get(
            reverse("admin:content_documenttranslation_change", args=[self.dt.pk])
        )
        self.assertContains(response, f"/export/{self.dt.pk}/docx/")
