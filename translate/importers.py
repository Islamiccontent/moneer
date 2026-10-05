"""مساعدات استيراد بيانات القرآن المعتمدة من ملفات xlsx: قراءة الورقة وتنظيف نص الترجمة."""

import re
from pathlib import Path

from django.core.management.base import CommandError
from openpyxl import load_workbook

_LEADING_NUMBER = re.compile(r"^\s*\d+\s*[.),:،]?\s*")
_BRACKETED_NUMBER = re.compile(r"\*?\(\d+\)|\[\d+\]|\{\d+\}")
_TRAILING_DOT = re.compile(r"\s*\.\s*$")
_SPACES = re.compile(r"\s+")


def clean_translation(text):
    """يحذف أرقام الحواشي ويوحّد الفراغات والنقطة الختامية."""
    text = _LEADING_NUMBER.sub("", str(text))
    text = _BRACKETED_NUMBER.sub("", text)
    text = _TRAILING_DOT.sub(".", text)
    return _SPACES.sub(" ", text).strip()


def read_sheet(path, required_columns):
    """يعيد (columns, rows) من الورقة الأولى؛ يرفع CommandError عند تعذر القراءة أو غياب عمود."""
    path = Path(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise CommandError(f"تعذّر قراءة الملف {path}: {exc}") from exc
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise CommandError(f"الملف {path} فارغ.")
        columns = {
            str(name).strip().lower(): i for i, name in enumerate(header) if name is not None
        }
        missing = [name for name in required_columns if name not in columns]
        if missing:
            raise CommandError(f"الملف {path} ينقصه الأعمدة: {', '.join(missing)}.")
        return columns, [row for row in rows if any(cell is not None for cell in row)]
    finally:
        workbook.close()


def cell(row, columns, name):
    """قيمة الخلية باسم عمودها، مشذَّبةً نصاً، أو سلسلة فارغة."""
    index = columns[name]
    value = row[index] if index < len(row) else None
    return "" if value is None else str(value).strip()


def to_int(value):
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return None
