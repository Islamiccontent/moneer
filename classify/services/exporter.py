"""تصدير المقاطع إلى جدول Excel بصيغة خط الإنتاج وقواعد ترقيمه المعتمدة."""

import uuid
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from . import classifier
from .taxonomy import (
    CONTENT_TYPE_NAME_AR,
    CONTENT_TYPE_OF_KIND,
    KIND_OF_INTERNAL,
    REASON_CODE,
    Align,
    ConfidenceReason,
    Direction,
    Tag,
    TextType,
)

COLUMNS = [
    "التسلسل",
    "رقم الفقرة",
    "الترتيب بالفقرة",
    "المحتوى",
    "نوع المحتوى",
    "التوسيم",
    "اتجاه النص",
    "محاذاة النص",
    "رقم الاصدار",
    "رقم الصفحة",
    "الحالة",
    "رقم الجزء",
    "csv_id",
    "نوع النص",
    "رقم الهامش",
    "رقم الأب للهامش",
    "uuid",
    "قابل للترجمة",
    "الاصدار-تلقائي",
    "رقم مرجعي",
    "نقحرة-نص",
    "الثقة",
    "سبب الثقة",
    "المصدر",
]

CONFIDENCE = {
    "title": (1.0, "نمط عنوان"),
    "attribution": (0.6, "صيغة إسناد"),
    "citation": (1.0, "إحالة مصدر"),
    "term": (0.6, "مدخل معجم"),
    "text": (0.4, "لا قرينة"),
    "athar": (0.6, "أثر — نُسب بصيغة التقديم"),
    "uncertain": (0.2, "لم تُحسم المطابقة"),
}
LOW_CONFIDENCE = {classifier.UNMATCHED_QURAN: 0.2, classifier.UNMATCHED_HADITH: 0.25}


def _score(s) -> float | None:
    """الدرجة من المصدرين: بنيةُ المصنِّف تسمّيها score، والنموذجُ match_score."""
    v = getattr(s, "score", None)
    return v if v is not None else getattr(s, "match_score", None)


MUSHAF_VARIANT = "يخالف رسم المصحف"


def confidence(s) -> tuple[float, str]:
    score = _score(s)
    if s.kind == "quran" and MUSHAF_VARIANT in (s.note or ""):
        return 0.2, "⚠ " + s.note.split("⚠ ", 1)[-1]
    if s.kind in ("quran", "hadith", "athar") and score is not None:
        return round(score, 3), f"مطابقة {int(score * 100)}%"
    if s.note == "إحالة مصدر":
        return 1.0, "إحالة مصدر"
    if s.note in LOW_CONFIDENCE:
        return LOW_CONFIDENCE[s.note], s.note
    return CONFIDENCE.get(s.kind, (0.4, "لا قرينة"))


def reason_code(s) -> ConfidenceReason:
    """رمزُ سبب الثقة للمرحلة التالية. نصُّه العربيّ يبقى للمراجع."""
    _, why = confidence(s)
    if s.kind == "quran" and MUSHAF_VARIANT in (s.note or ""):
        return ConfidenceReason.UNMATCHED_QURAN
    if s.kind in ("quran", "hadith") or (s.kind == "athar" and _score(s) is not None):
        return ConfidenceReason.INDEX_MATCH
    if s.note == classifier.UNMATCHED_QURAN:
        return ConfidenceReason.UNMATCHED_QURAN
    if s.note == classifier.UNMATCHED_HADITH:
        return ConfidenceReason.UNMATCHED_HADITH
    return REASON_CODE.get(why, ConfidenceReason.NO_EVIDENCE)


CONTENT_TYPE_ID = {
    internal: CONTENT_TYPE_OF_KIND[kind] for internal, kind in KIND_OF_INTERNAL.items()
}
CONTENT_TYPE = {k: CONTENT_TYPE_NAME_AR[v] for k, v in CONTENT_TYPE_ID.items()}


def build_rows(segments) -> list[dict]:
    """بناء صفوف الجدول من مقاطع مستند واحد."""
    by_para: dict[int, list] = {}
    for s in segments:
        by_para.setdefault(s.para or 0, []).append(s)

    rows, seq, para_no = [], 1, 1
    for key in sorted(by_para):
        group = by_para[key]
        for i, s in enumerate(group, 1):
            heading = s.kind == "title" and s.level
            footnote = getattr(s, "footnote", False)
            order = i
            conf, why = confidence(s)
            csv_id = f"{para_no * 10}" + ("" if order == 1 else f"_{(order - 1) * 10}")
            rows.append(
                {
                    "التسلسل": seq,
                    "رقم الفقرة": para_no,
                    "الترتيب بالفقرة": order,
                    "المحتوى": s.content,
                    "نوع المحتوى": "هامش" if footnote else CONTENT_TYPE.get(s.kind, "نص"),
                    "التوسيم": (
                        Tag.FOOTNOTE
                        if footnote
                        else Tag[f"HEADING_{min(s.level, 5)}"]
                        if heading
                        else Tag.PARAGRAPH
                    ).value,
                    "اتجاه النص": Direction.RTL.value,
                    "محاذاة النص": (Align.CENTER if heading else Align.RIGHT).value,
                    "رقم الاصدار": 1,
                    "رقم الصفحة": 1,
                    "الحالة": "new",
                    "رقم الجزء": 1,
                    "csv_id": csv_id,
                    "نوع النص": (TextType.FOOTNOTE if footnote else TextType.PARAGRAPH).value,
                    "رقم الهامش": "",
                    "رقم الأب للهامش": "",
                    "uuid": str(uuid.uuid4()),
                    "قابل للترجمة": True,
                    "الاصدار-تلقائي": 1,
                    "رقم مرجعي": "",
                    "نقحرة-نص": "p",
                    "الثقة": conf,
                    "سبب الثقة": why,
                    "المصدر": s.source,
                }
            )
            seq += 1
        para_no += 1
    return rows


def to_xlsx(title: str, segments) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "".join(c for c in title if c not in "[]:*?/\\")[:31] or "Sheet1"
    ws.sheet_view.rightToLeft = True

    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in build_rows(segments):
        ws.append([row[c] for c in COLUMNS])

    widths = {
        "المحتوى": 70,
        "uuid": 38,
        "csv_id": 12,
        "نوع المحتوى": 12,
        "التوسيم": 11,
        "سبب الثقة": 16,
        "المصدر": 26,
        "الثقة": 8,
    }
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 14)
    content_col = COLUMNS.index("المحتوى") + 1
    for row in ws.iter_rows(min_row=2, min_col=content_col, max_col=content_col):
        row[0].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
