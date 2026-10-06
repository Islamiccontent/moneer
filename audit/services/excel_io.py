"""كتابة ملف نتائج التدقيق بأوراقه الأربع (منقول من trreview/excel_io.py بلا قراءة ملف الإدخال).

الإدخال هنا جداول تطبيق audit لا ملف Excel؛ ``write_output`` و``compute_stats`` تأخذان
قواميس بالأعمدة نفسها التي كانت تُقرأ من review.db (انظر ``audit.pipeline.export_rows``).

الإخراج: أربع أوراق —
    الترجمات    صف لكل ترجمة: الأعمدة الأصلية + الحكم + ملخص الأخطاء
    الأخطاء     صف لكل خطأ مصنَّف (الفئة، النوع، الرمز، الخطورة، ...)
    إحصائيات    مؤشرات على مستوى المهمة (حسب الفئة/النوع/الخطورة/الكشف)
    التصنيف     مرجع التصنيف المعتمد
"""

import json
from collections import Counter, OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .taxonomy import CATEGORIES, DETECTIONS, SEVERITIES, VERDICTS, taxonomy_rows

# ══════════════════════ الكتابة ══════════════════════

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial", size=11)
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

VERDICT_FILL = {
    "مقبول": PatternFill("solid", fgColor="E2F0D9"),
    "يحتاج مراجعة": PatternFill("solid", fgColor="FFF2CC"),
    "مرفوض": PatternFill("solid", fgColor="F8CBAD"),
    "لم يُفحص": PatternFill("solid", fgColor="EDEDED"),
}
SEVERITY_FILL = {
    "حرج": PatternFill("solid", fgColor="C00000"),
    "كبير": PatternFill("solid", fgColor="F4B183"),
    "متوسط": PatternFill("solid", fgColor="FFE699"),
    "بسيط": PatternFill("solid", fgColor="DDEBF7"),
}
SEVERITY_FONT = {"حرج": Font(name="Arial", size=10, bold=True, color="FFFFFF")}


def _write_table(ws, headers, rows, widths=None, rtl=True, freeze="A2"):
    ws.sheet_view.rightToLeft = rtl
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font, cell.alignment, cell.border = HEADER_FILL, HEADER_FONT, CENTER, BORDER
    for row in rows:
        ws.append(row)
    for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in r:
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP_TOP, BORDER
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(header, 18)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30
    if freeze:
        ws.freeze_panes = freeze
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def _sev_text(s):
    return s or ""


def build_translations_sheet(wb, job, rows, findings_by_row):
    ws = wb.active
    ws.title = "الترجمات"
    columns = json.loads(job["columns_json"])
    extra_headers = [
        c for c in columns["headers"] if c not in (columns["source_col"], columns["target_col"])
    ]
    headers = [
        "#",
        "رقم الصف في الأصل",
        columns["source_col"],
        columns["target_col"],
        *extra_headers,
        "الحكم",
        "عدد الأخطاء",
        "أعلى خطورة",
        "حساسية شرعية",
        "يحتاج مراجعة بشرية",
        "الفئات",
        "أنواع الأخطاء",
        "ملخص المراجعة",
        "حالة الفحص",
    ]
    data = []
    for i, r in enumerate(rows, start=1):
        extra = json.loads(r["extra_json"] or "{}")
        fl = findings_by_row.get(r["id"], [])
        verdict = r["verdict"] or ("لم يُفحص" if r["status"] != "done" else "مقبول")
        cats = " ، ".join(OrderedDict.fromkeys(f["category"] for f in fl))
        types = "\n".join(f"{f['type_code']} — {f['type_name']}" for f in fl)
        status_label = {
            "done": "تم",
            "pending": "لم يُفحص بعد",
            "failed": "فشل الفحص",
            "skipped": "مستبعد",
        }.get(r["status"], r["status"])
        if r["skip_reason"]:
            status_label += f" ({r['skip_reason']})"
        data.append(
            [
                i,
                r["excel_row"],
                r["source_text"],
                r["target_text"],
                *[extra.get(h) for h in extra_headers],
                verdict,
                len(fl),
                _sev_text(r["max_severity"]),
                "نعم" if any(f["sharia_sensitive"] for f in fl) else "",
                "نعم" if r["needs_human_review"] else "",
                cats,
                types,
                r["summary"] or "",
                status_label,
            ]
        )
    widths = {
        "#": 6,
        "رقم الصف في الأصل": 9,
        columns["source_col"]: 55,
        columns["target_col"]: 55,
        "الحكم": 14,
        "عدد الأخطاء": 9,
        "أعلى خطورة": 11,
        "حساسية شرعية": 10,
        "يحتاج مراجعة بشرية": 11,
        "الفئات": 24,
        "أنواع الأخطاء": 36,
        "ملخص المراجعة": 45,
        "حالة الفحص": 16,
    }
    for h in extra_headers:
        widths.setdefault(h, 14)
    _write_table(ws, headers, data, widths, freeze="E2")
    verdict_col = headers.index("الحكم") + 1
    sev_col = headers.index("أعلى خطورة") + 1
    for row_idx in range(2, ws.max_row + 1):
        v = ws.cell(row=row_idx, column=verdict_col)
        if v.value in VERDICT_FILL:
            v.fill = VERDICT_FILL[v.value]
        s = ws.cell(row=row_idx, column=sev_col)
        if s.value in SEVERITY_FILL:
            s.fill = SEVERITY_FILL[s.value]
            if s.value in SEVERITY_FONT:
                s.font = SEVERITY_FONT[s.value]
    return ws


def build_findings_sheet(wb, findings):
    ws = wb.create_sheet("الأخطاء")
    headers = [
        "#",
        "رقم الصف في الأصل",
        "النص الأصلي",
        "الترجمة",
        "الفئة",
        "رمز الفئة",
        "نوع الخطأ",
        "رمز النوع",
        "الخطورة",
        "حساسية شرعية",
        "طريقة الكشف",
        "مصدر الكشف",
        "الثقة %",
        "موضع الخطأ",
        "شرح المشكلة",
        "المطلوب",
        "التصحيح المقترح",
        "يحتاج مراجعة بشرية",
    ]
    data = []
    for i, f in enumerate(findings, start=1):
        data.append(
            [
                i,
                f["excel_row"],
                f["source_text"],
                f["target_text"],
                f["category"],
                f["category_code"],
                f["type_name"],
                f["type_code"],
                f["severity"],
                "نعم" if f["sharia_sensitive"] else "لا",
                f["detection"],
                "قاعدة آلية" if f["source"] == "rule" else "النموذج",
                f["confidence"],
                f["location"],
                f["explanation"],
                f["required_action"],
                f["suggested_fix"],
                "نعم" if f["needs_human_review"] else "",
            ]
        )
    widths = {
        "#": 6,
        "رقم الصف في الأصل": 9,
        "النص الأصلي": 45,
        "الترجمة": 45,
        "الفئة": 22,
        "رمز الفئة": 8,
        "نوع الخطأ": 28,
        "رمز النوع": 11,
        "الخطورة": 9,
        "حساسية شرعية": 9,
        "طريقة الكشف": 12,
        "مصدر الكشف": 11,
        "الثقة %": 7,
        "موضع الخطأ": 30,
        "شرح المشكلة": 45,
        "المطلوب": 26,
        "التصحيح المقترح": 40,
        "يحتاج مراجعة بشرية": 10,
    }
    _write_table(ws, headers, data, widths, freeze="E2")
    sev_col = headers.index("الخطورة") + 1
    for row_idx in range(2, ws.max_row + 1):
        s = ws.cell(row=row_idx, column=sev_col)
        if s.value in SEVERITY_FILL:
            s.fill = SEVERITY_FILL[s.value]
            if s.value in SEVERITY_FONT:
                s.font = SEVERITY_FONT[s.value]
    return ws


def compute_stats(job, rows, findings):
    """كل المؤشرات المستخدمة في ورقة الإحصائيات — دالة مستقلة كي تُختبر."""
    total = len(rows)
    checked = sum(1 for r in rows if r["status"] == "done")
    skipped = sum(1 for r in rows if r["status"] == "skipped")
    failed = sum(1 for r in rows if r["status"] == "failed")
    pending = total - checked - skipped - failed
    with_errors = sum(1 for r in rows if (r["error_count"] or 0) > 0)
    source_words = sum(
        len(str(r["source_text"]).split()) for r in rows if r["status"] in ("done", "skipped")
    )
    n_errors = len(findings)
    considered = checked + skipped

    def pct(part, whole):
        return round(100.0 * part / whole, 1) if whole else 0.0

    verdicts = Counter(r["verdict"] for r in rows if r["verdict"])
    by_cat = Counter(f["category"] for f in findings)
    by_type = Counter((f["type_code"], f["type_name"]) for f in findings)
    by_sev = Counter(f["severity"] for f in findings)
    by_det = Counter(f["detection"] for f in findings)
    by_src = Counter(f["source"] for f in findings)
    sharia = sum(1 for f in findings if f["sharia_sensitive"])
    human = sum(1 for r in rows if r["needs_human_review"])
    cat_order = [c.name for c in CATEGORIES.values()] + ["غير مصنَّف"]

    return {
        "general": [
            ("اسم الملف", job["input_file"]),
            ("اللغة", job["language"]),
            ("النموذج", job["model"]),
            ("نسخة البرومت", job["prompt_version"]),
            ("تاريخ المهمة", job["created_at"]),
            ("إجمالي الصفوف", total),
            ("صفوف فُحصت بالنموذج", checked),
            ("صفوف مستبعدة بقاعدة قطعية (فارغة/نسخ)", skipped),
            ("صفوف فشل فحصها", failed),
            ("صفوف لم تُفحص بعد", pending),
            ("صفوف فيها أخطاء", with_errors),
            ("نسبة الصفوف التي فيها أخطاء %", pct(with_errors, considered)),
            ("إجمالي الأخطاء", n_errors),
            ("أخطاء لكل 100 صف", round(100.0 * n_errors / considered, 2) if considered else 0),
            (
                "أخطاء لكل 1000 كلمة (من الأصل)",
                round(1000.0 * n_errors / source_words, 2) if source_words else 0,
            ),
            ("أخطاء ذات حساسية شرعية", sharia),
            ("صفوف تحتاج مراجعة بشرية إلزامية", human),
            ("توكن الإدخال (تقريبي من Gemini)", job["prompt_tokens"]),
            ("توكن الإخراج (تقريبي من Gemini)", job["output_tokens"]),
        ],
        "verdicts": [
            (v, verdicts.get(v, 0), pct(verdicts.get(v, 0), considered)) for v in VERDICTS
        ],
        "by_category": [
            (c, by_cat.get(c, 0), pct(by_cat.get(c, 0), n_errors))
            for c in cat_order
            if by_cat.get(c, 0) or c != "غير مصنَّف"
        ],
        "by_type": [(code, name, n, pct(n, n_errors)) for (code, name), n in by_type.most_common()],
        "by_severity": [(s, by_sev.get(s, 0), pct(by_sev.get(s, 0), n_errors)) for s in SEVERITIES],
        "by_detection": [
            (d, by_det.get(d, 0), pct(by_det.get(d, 0), n_errors)) for d in DETECTIONS
        ],
        "by_source": [
            ("قاعدة آلية", by_src.get("rule", 0), pct(by_src.get("rule", 0), n_errors)),
            ("النموذج", by_src.get("model", 0), pct(by_src.get("model", 0), n_errors)),
        ],
    }


def build_stats_sheet(wb, stats):
    ws = wb.create_sheet("إحصائيات")
    ws.sheet_view.rightToLeft = True
    title_font = Font(name="Arial", size=12, bold=True, color="1F4E5F")
    row = 1

    def section(title, headers, data):
        nonlocal row
        ws.cell(row=row, column=1, value=title).font = title_font
        row += 1
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.fill, cell.font, cell.alignment, cell.border = (
                HEADER_FILL,
                HEADER_FONT,
                CENTER,
                BORDER,
            )
        row += 1
        for values in data:
            for c, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=c, value=v)
                cell.font, cell.border, cell.alignment = BODY_FONT, BORDER, WRAP_TOP
            row += 1
        row += 1

    section("عام", ["المؤشر", "القيمة"], stats["general"])
    section("توزيع الأحكام", ["الحكم", "العدد", "النسبة %"], stats["verdicts"])
    section("الأخطاء حسب الخطورة", ["الخطورة", "العدد", "النسبة %"], stats["by_severity"])
    section("الأخطاء حسب الفئة", ["الفئة", "العدد", "النسبة %"], stats["by_category"])
    section("الأخطاء حسب طريقة الكشف", ["طريقة الكشف", "العدد", "النسبة %"], stats["by_detection"])
    section("الأخطاء حسب مصدر الكشف", ["المصدر", "العدد", "النسبة %"], stats["by_source"])
    section(
        "الأخطاء حسب النوع (الأكثر تكرارًا أولًا)",
        ["الرمز", "نوع الخطأ", "العدد", "النسبة %"],
        stats["by_type"],
    )
    for col, width in zip("ABCD", (38, 34, 12, 12), strict=True):
        ws.column_dimensions[col].width = width
    return ws


def build_taxonomy_sheet(wb):
    ws = wb.create_sheet("التصنيف")
    rows = taxonomy_rows()
    headers = list(rows[0].keys())
    widths = {
        "رمز الفئة": 8,
        "الفئة": 24,
        "رمز النوع": 10,
        "نوع الخطأ": 30,
        "التعريف": 60,
        "الخطورة الافتراضية": 10,
        "حساسية شرعية افتراضية": 10,
        "طريقة الكشف الافتراضية": 12,
        "المطلوب": 30,
    }
    _write_table(ws, headers, [[r[h] for h in headers] for r in rows], widths)
    return ws


def write_output(path, job, rows, findings):
    """يكتب ملف النتائج. ``rows`` صفوف المهمة، ``findings`` ملاحظاتها (مع excel_row)."""
    findings_by_row = {}
    for f in findings:
        findings_by_row.setdefault(f["row_id"], []).append(f)
    wb = Workbook()
    build_translations_sheet(wb, job, rows, findings_by_row)
    build_findings_sheet(wb, findings)
    build_stats_sheet(wb, compute_stats(job, rows, findings))
    build_taxonomy_sheet(wb)
    wb.save(path)
    return path
