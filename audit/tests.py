"""اختبارات أداة التدقيق: القواعد والتصنيف والتحليل والبرومت، والخط كاملاً بموصل Gemini مزيَّف،
والمهمة الخلفية والواجهة والأمر وadmin."""

import json
import tempfile
import uuid
from io import BytesIO, StringIO
from pathlib import Path
from unittest import mock
from urllib.parse import unquote

from django.conf import settings as dj_settings
from django.contrib import admin
from django.core.management import call_command
from django.core.management.base import CommandError
from django.tasks import TaskResultStatus
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from content.models import Document, DocumentTranslation, Phrase, PhraseTranslation
from core.models import ContentTypeLookup, Language
from users.models import User

from . import pipeline, tasks
from .models import AuditFinding, AuditJob, AuditRow
from .services import gemini, local_rules, parsing, prompt, taxonomy
from .services.config import load_settings
from .services.languages import is_arabic_script
from .views import XLSX_TYPE

IMMEDIATE_TASKS = {"default": {"BACKEND": "django.tasks.backends.immediate.ImmediateBackend"}}
SMALL_AUDIT = {**dj_settings.AUDIT, "GROUP_SIZE": 5, "MAX_ATTEMPTS": 2}


class FakeGeminiClient:
    """يحاكي Gemini فورياً: يعيد رداً مصنَّفاً لكل عنصر في البرومت (منقول من اختبارات المنصة).

    ``error_every``: كل عنصر رقمه من مضاعفاته يحصل على SEM-OMI، ومضاعفات ثلاثة أضعافه على
    SEM-NEG بحساسية شرعية. ``broken_calls``: أرقام النداءات (من 1) التي تعيد رداً غير صالح.
    ``raise_always``: كل نداء يرفع استثناءً.
    """

    def __init__(self, settings=None, error_every=3, broken_calls=(), raise_always=False):
        self.model = getattr(settings, "model", "fake-model")
        self.error_every = error_every
        self.broken_calls = set(broken_calls)
        self.raise_always = raise_always
        self.calls = []

    def build_request(self, prompt_text):
        return {"contents": [{"role": "user", "parts": [{"text": prompt_text}]}]}

    def _answer(self, prompt_text):
        ids = [
            int(line.split("id=")[1].split(" ")[0].strip())
            for line in prompt_text.splitlines()
            if line.startswith("--- عنصر id=")
        ]
        answer = []
        for i in ids:
            errors = []
            if i % self.error_every == 0:
                errors.append(
                    {
                        "code": "SEM-OMI",
                        "severity": "كبير",
                        "sharia_sensitive": False,
                        "detection": "استدلالي",
                        "confidence": 88,
                        "location": "…",
                        "explanation": "سقطت جملة من الترجمة.",
                        "required_action": "إضافة المعنى المفقود",
                        "suggested_fix": "…",
                    }
                )
            if i % (self.error_every * 3) == 0:
                errors.append(
                    {
                        "code": "SEM-NEG",
                        "severity": "حرج",
                        "sharia_sensitive": True,
                        "detection": "استدلالي",
                        "confidence": 95,
                        "location": "…",
                        "explanation": "نفي قُلب إثباتًا.",
                        "suggested_fix": "…",
                    }
                )
            answer.append({"id": i, "summary": "فيها ملاحظات." if errors else "", "errors": errors})
        return json.dumps(answer, ensure_ascii=False)

    def generate(self, request_body):
        self.calls.append(request_body)
        if self.raise_always:
            raise gemini.GeminiAPIError(503, "خطأ مزيَّف")
        if len(self.calls) in self.broken_calls:
            text = "ليس JSON"
        else:
            text = (
                "```json\n"
                + self._answer(request_body["contents"][0]["parts"][0]["text"])
                + "\n```"
            )
        return {
            "response": {
                "candidates": [{"content": {"parts": [{"text": text}]}}],
                "usageMetadata": {"promptTokenCount": 1000, "candidatesTokenCount": 200},
            }
        }


def fake_gemini(**kwargs):
    """يستبدل GeminiClient داخل الخط بموصل مزيَّف واحد مشترك بين النداءات؛ يعيد (patch, shared)."""
    shared = {}

    def factory(settings):
        if "client" not in shared:
            shared["client"] = FakeGeminiClient(settings, **kwargs)
        return shared["client"]

    return mock.patch.object(pipeline, "GeminiClient", side_effect=factory), shared


# ══════════════════════ الوحدات المنقولة كما هي ══════════════════════


class LocalRulesTests(SimpleTestCase):
    """القواعد القطعية ومؤشرات الاشتباه تعمل قبل النموذج وبلا تكلفة."""

    def rows(self, *targets, source="هذا نص عربي طويل فيه كلمات كثيرة للتجربة والفحص"):
        return [
            {"source": f"{source} {i}", "target": t, "row_type": "translation"}
            for i, t in enumerate(targets)
        ]

    def test_empty_and_copied_translations_are_skipped_with_definitive_findings(self):
        results = local_rules.apply_rules(
            self.rows("", "   ", "هذا نص عربي طويل فيه كلمات كثيرة للتجربة والفحص 2")
        )
        self.assertEqual(
            [r["skip"] for r in results], ["ترجمة فارغة", "ترجمة فارغة", "نسخ النص الأصلي"]
        )
        self.assertEqual(
            [r["findings"][0]["code"] for r in results], ["FRM-EMP", "FRM-EMP", "FRM-CPY"]
        )

    def test_arabic_letters_are_flagged_unless_the_language_uses_arabic_script(self):
        rows = self.rows(
            "Text with كلمة inside", "Text with ﷺ and ۝ only", "Text with صلى الله عليه وسلم"
        )
        codes = [[f["code"] for f in r["findings"]] for r in local_rules.apply_rules(rows)]
        self.assertEqual(codes, [["FRM-ARB"], [], []])
        self.assertEqual(
            [r["findings"] for r in local_rules.apply_rules(rows, arabic_script=True)], [[], [], []]
        )

    def test_unbalanced_brackets_only_when_the_source_is_balanced(self):
        rows = [
            {"source": "نص (سليم) هنا", "target": "Text (broken here"},
            {"source": "نص (مكسور هنا", "target": "Text (broken too"},
        ]
        codes = [[f["code"] for f in r["findings"]] for r in local_rules.apply_rules(rows)]
        self.assertEqual(codes, [["FRM-BRK"], []])

    def test_duplicate_target_for_different_sources_becomes_a_hint_not_a_finding(self):
        rows = [
            {"source": "الحمد لله رب العالمين على نعمه الكثيرة", "target": "Same translation here"},
            {
                "source": "الصلاة عماد الدين ومن تركها فقد هدم الدين",
                "target": "Same translation here",
            },
        ]
        results = local_rules.apply_rules(rows)
        self.assertEqual([r["findings"] for r in results], [[], []])
        self.assertTrue(all("الترجمة نفسها" in r["hints"][0] for r in results))

    def test_arabic_script_languages(self):
        self.assertTrue(is_arabic_script("ur"))
        self.assertTrue(is_arabic_script(" FA "))
        self.assertFalse(is_arabic_script("en"))
        self.assertFalse(is_arabic_script(""))


class TaxonomyTests(SimpleTestCase):
    def test_verdict_follows_the_highest_severity(self):
        self.assertEqual(taxonomy.verdict_for([]), "مقبول")
        self.assertEqual(taxonomy.verdict_for(["بسيط", "متوسط"]), "يحتاج مراجعة")
        self.assertEqual(taxonomy.verdict_for(["بسيط", "كبير"]), "مرفوض")
        self.assertEqual(taxonomy.max_severity(["متوسط", "حرج", "بسيط"]), "حرج")
        self.assertIsNone(taxonomy.max_severity([]))

    def test_unknown_codes_fall_back_to_other_and_aliases_are_normalized(self):
        self.assertIs(taxonomy.get_type("XX-YY"), taxonomy.OTHER_TYPE)
        self.assertEqual(taxonomy.get_type("sem_omi").code, "SEM-OMI")
        self.assertEqual(taxonomy.normalize_severity("major", "بسيط"), "كبير")
        self.assertEqual(taxonomy.normalize_severity("غريب", "بسيط"), "بسيط")
        self.assertEqual(taxonomy.normalize_detection("rule", "استدلالي"), "قطعي")

    def test_prompt_table_lists_every_type_once(self):
        text = taxonomy.taxonomy_table_text()
        for code in taxonomy.ERROR_TYPES:
            self.assertEqual(text.count(f"- {code} |"), 1, code)
        self.assertEqual(len(taxonomy.taxonomy_rows()), len(taxonomy.ERROR_TYPES))


class ParsingTests(SimpleTestCase):
    def test_json_array_is_extracted_from_fences_or_wrappers(self):
        self.assertEqual(parsing.parse_json_array('```json\n[{"id": 1}]\n```'), [{"id": 1}])
        self.assertEqual(parsing.parse_json_array('{"items": [{"id": 2}]}'), [{"id": 2}])
        self.assertEqual(parsing.parse_json_array('نص زائد [{"id": 3}] وبعده'), [{"id": 3}])
        with self.assertRaises(ValueError):
            parsing.parse_json_array("لا شيء هنا")

    def test_normalize_finding_derives_category_and_defaults_from_the_code(self):
        finding = parsing.normalize_finding({"code": "SHR-NUM", "confidence": "0.9"})
        self.assertEqual(
            (finding["category_code"], finding["type_name"]), ("SHR", "رقم شرعي مغلوط")
        )
        self.assertEqual((finding["severity"], finding["detection"]), ("حرج", "قطعي"))
        self.assertTrue(finding["sharia_sensitive"] and finding["needs_human_review"])
        self.assertEqual(finding["confidence"], 90)
        self.assertEqual(finding["required_action"], "تصحيح الرقم من الأصل")
        unknown = parsing.normalize_finding({"code": "ZZZ"}, source="rule")
        self.assertEqual((unknown["type_code"], unknown["type_name"]), ("OTHER", "غير مصنَّف (ZZZ)"))
        self.assertEqual(unknown["confidence"], 100)

    def test_group_result_requires_every_expected_id_and_dedupes_errors(self):
        duplicated = [{"code": "SEM-OMI", "location": "x"}] * 2
        text = json.dumps([{"id": 1, "summary": "", "errors": duplicated}, {"id": 2, "errors": []}])
        parsed = parsing.validate_group_result(text, [1, 2])
        self.assertEqual(len(parsed[1]["errors"]), 1)
        self.assertEqual(parsed[2]["errors"], [])
        with self.assertRaisesMessage(ValueError, "رد ناقص"):
            parsing.validate_group_result(text, [1, 2, 3])


class PromptTests(SimpleTestCase):
    def test_prompt_includes_items_count_and_language_rules(self):
        items = [
            {"id": 7, "source": "نص", "target": "Text", "row_type": "naqhara", "hints": ["تنبيه"]},
            {"id": 9, "source": "نص آخر", "target": "Other", "row_type": "translation"},
        ]
        text = prompt.build_prompt(items, "الإنجليزية", extra_rules="لا تترجم الأعلام")
        self.assertIn("--- عنصر id=7 ---", text)
        self.assertIn("النوع المطلوب: نقحرة", text)
        self.assertIn("تنبيهات آلية: تنبيه", text)
        self.assertIn("وعدد عناصرها 2 بالضبط", text)
        self.assertIn("لغة الإنجليزية لا تُكتب بحروف عربية", text)
        self.assertIn("## قواعد إضافية خاصة بهذه اللغة/المشروع\nلا تترجم الأعلام", text)
        urdu = prompt.build_prompt(items, "الأردية", arabic_script=True)
        self.assertIn("تُكتب أصلًا بحروف عربية", urdu)
        self.assertNotIn("{language}", urdu)


class GeminiModuleTests(SimpleTestCase):
    def test_request_and_response_helpers(self):
        body = gemini.build_request("p", "high", 0.2, {"type": "ARRAY"})
        self.assertEqual(body["generationConfig"]["thinkingConfig"], {"thinkingLevel": "high"})
        self.assertEqual(body["generationConfig"]["responseSchema"], {"type": "ARRAY"})
        item = {
            "response": {
                "candidates": [
                    {"content": {"parts": [{"text": "فكرة", "thought": True}, {"text": "[]"}]}}
                ],
                "usageMetadata": {"promptTokenCount": 5, "candidatesTokenCount": 2},
            }
        }
        self.assertEqual(gemini.response_text(item), "[]")
        self.assertEqual(gemini.usage_tokens(item), (5, 2))
        with self.assertRaisesMessage(ValueError, "blockReason=SAFETY"):
            gemini.response_text({"response": {"promptFeedback": {"blockReason": "SAFETY"}}})
        with self.assertRaisesMessage(ValueError, "خطأ"):
            gemini.response_text({"error": {"message": "خطأ"}})
        self.assertTrue(gemini.GeminiAPIError(404, "x").permanent)
        self.assertFalse(gemini.GeminiAPIError(503, "x").permanent)


# ══════════════════════ الخط على جداول المشروع ══════════════════════


@override_settings(GEMINI_API_KEY="test-key", AUDIT=SMALL_AUDIT, TASKS=IMMEDIATE_TASKS)
class AuditFixture(TestCase):
    """مستند باثنتي عشرة جملة قابلة للترجمة (إحداها نقحرة) وجملة غير قابلة، وترجمة إنجليزية فيها
    جملة بلا ترجمة (FRM-EMP) وأخرى منسوخة (FRM-CPY) وثالثة فيها حرف عربي (FRM-ARB)."""

    @classmethod
    def setUpTestData(cls):
        cls.english = Language.objects.get(iso_code="en")
        cls.document = Document.objects.create(title="كتاب التوحيد", kind=Document.Kind.ARTICLE)
        text_type = ContentTypeLookup.objects.get(code="text")
        cls.phrases = [
            Phrase.objects.create(
                document=cls.document,
                content_type=text_type,
                group_id=i,
                group_order=1,
                text=f"هذا نص عربي رقم {i} فيه كلمات كثيرة للتجربة والفحص والمراجعة",
                text_type=Phrase.TextType.NAQHARA if i == 11 else Phrase.TextType.P,
            )
            for i in range(1, 13)
        ]
        cls.untranslatable = Phrase.objects.create(
            document=cls.document,
            content_type=text_type,
            group_id=13,
            group_order=1,
            text="[البقرة: 153]",
            translatable=False,
        )
        cls.dt = DocumentTranslation.objects.create(
            document=cls.document,
            target_language=cls.english,
            status=DocumentTranslation.Status.REVIEW,
        )
        for i, phrase in enumerate(cls.phrases, start=1):
            if i == 2:
                continue
            if i == 4:
                text = phrase.text
            elif i == 7:
                text = f"This is an English text number {i} with the word كلمة inside it"
            else:
                text = f"This is an English text number {i} with many words for the test"
            PhraseTranslation.objects.create(
                document_translation=cls.dt,
                phrase=phrase,
                ai_translation=text,
                translation=text,
                status=PhraseTranslation.Status.SUGGESTED,
            )


class CreateJobTests(AuditFixture):
    def test_rows_rules_and_groups_are_built_without_calling_the_model(self):
        with mock.patch.object(pipeline.gemini, "generate_content") as generate:
            job = pipeline.create_job(self.dt)
        generate.assert_not_called()
        self.assertEqual(job.status, AuditJob.Status.PENDING)
        self.assertEqual(
            (job.language, job.arabic_script, job.model),
            ("الإنجليزية", False, SMALL_AUDIT["MODEL"]),
        )
        self.assertEqual((job.group_size, job.prompt_version), (5, prompt.PROMPT_VERSION))
        self.assertEqual(job.prompt_snapshot, prompt.PROMPT_TEMPLATE)
        self.assertEqual(job.request_config["thinking_level"], SMALL_AUDIT["THINKING_LEVEL"])
        self.assertEqual((job.total_rows, job.eligible_rows), (12, 10))
        self.assertEqual(job.rows.count(), 12)
        self.assertFalse(job.rows.filter(phrase=self.untranslatable).exists())

        empty, copied, arabic = job.rows.get(seq=2), job.rows.get(seq=4), job.rows.get(seq=7)
        self.assertEqual(
            (empty.status, empty.skip_reason, empty.verdict), ("skipped", "ترجمة فارغة", "مرفوض")
        )
        self.assertEqual(empty.findings.get().type_code, "FRM-EMP")
        self.assertEqual((copied.status, copied.findings.get().type_code), ("skipped", "FRM-CPY"))
        self.assertEqual((arabic.status, arabic.findings.get().type_code), ("pending", "FRM-ARB"))
        self.assertEqual(arabic.findings.get().source, AuditFinding.Source.RULE)
        self.assertEqual(job.rows.get(seq=11).row_type, AuditRow.RowType.NAQHARA)

        groups = list(job.groups.order_by("id"))
        self.assertEqual([g.key for g in groups], [f"j{job.pk}-g00000", f"j{job.pk}-g00001"])
        self.assertEqual([g.rows.count() for g in groups], [5, 5])
        self.assertEqual(list(groups[0].rows.values_list("seq", flat=True)), [1, 3, 5, 6, 7])
        self.assertIsNone(empty.group)

    def test_limit_restricts_the_rows_and_empty_translation_is_rejected(self):
        job = pipeline.create_job(self.dt, limit=3)
        self.assertEqual(list(job.rows.values_list("seq", flat=True)), [1, 2, 3])
        other = DocumentTranslation.objects.create(
            document=Document.objects.create(title="فارغ", kind=Document.Kind.POST),
            target_language=self.english,
            status=DocumentTranslation.Status.REVIEW,
        )
        with self.assertRaisesMessage(pipeline.AuditError, "لا توجد جمل"):
            pipeline.create_job(other)

    def test_group_prompt_uses_the_snapshot_and_the_row_hints(self):
        job = pipeline.create_job(self.dt, extra_rules="قاعدة خاصة")
        job.prompt_snapshot = "لغة {language} — {items} — {count} — {extra_rules}"
        text = pipeline.build_group_prompt(job, job.groups.first())
        self.assertTrue(text.startswith("لغة الإنجليزية — --- عنصر id=1 ---"))
        self.assertIn("— 5 —", text)
        self.assertIn("قاعدة خاصة", text)


class RunAuditTests(AuditFixture):
    def test_end_to_end_with_a_fake_model(self):
        patch, shared = fake_gemini()
        with patch:
            job = pipeline.run_audit(self.dt)
        self.assertEqual(job.status, AuditJob.Status.DONE)
        self.assertIsNotNone(job.finished_at)
        self.assertEqual(len(shared["client"].calls), 2)
        self.assertEqual((job.prompt_tokens, job.output_tokens), (2000, 400))
        self.assertEqual(set(job.groups.values_list("status", flat=True)), {"done"})
        self.assertTrue(all(g.raw_response.startswith("```json") for g in job.groups.all()))
        self.assertEqual(job.progress, {"groups": {"done": 2}, "rows": {"done": 10, "skipped": 2}})

        # SEM-OMI على مضاعفات 3 من المؤهلة (3، 6، 9، 12) وSEM-NEG على 9 بحساسية شرعية
        omitted = job.rows.filter(findings__type_code="SEM-OMI").values_list("seq", flat=True)
        self.assertEqual(sorted(omitted), [3, 6, 9, 12])
        row9 = job.rows.get(seq=9)
        self.assertEqual([f.type_code for f in row9.findings.all()], ["SEM-OMI", "SEM-NEG"])
        self.assertEqual((row9.verdict, row9.max_severity, row9.error_count), ("مرفوض", "حرج", 2))
        self.assertTrue(row9.sharia_sensitive and row9.needs_human_review)
        self.assertEqual(row9.summary, "فيها ملاحظات.")
        clean = job.rows.get(seq=5)
        self.assertEqual((clean.status, clean.verdict, clean.error_count), ("done", "مقبول", 0))
        # ملاحظة القاعدة على الصف 7 بقيت وترقيم ملاحظات النموذج يكمل بعدها
        self.assertEqual(
            list(job.rows.get(seq=7).findings.values_list("seq", "type_code")), [(1, "FRM-ARB")]
        )
        self.assertEqual(job.finding_count, 8)
        self.assertEqual(job.stats["general"]["إجمالي الصفوف"], 12)
        self.assertEqual(job.stats["general"]["صفوف فُحصت بالنموذج"], 10)
        self.assertEqual(job.stats["general"]["أخطاء ذات حساسية شرعية"], 1)
        verdicts = {v["label"]: v["count"] for v in job.stats["verdicts"]}
        self.assertEqual(verdicts, {"مقبول": 5, "يحتاج مراجعة": 1, "مرفوض": 6})
        self.assertEqual(
            job.stats["by_source"][0], {"label": "قاعدة آلية", "count": 3, "pct": 37.5}
        )
        self.assertEqual(job.stats["by_type"][0]["code"], "SEM-OMI")

    def test_broken_response_is_retried_in_the_same_run(self):
        patch, shared = fake_gemini(broken_calls={1})
        with patch:
            job = pipeline.run_audit(self.dt)
        self.assertEqual(job.status, AuditJob.Status.DONE)
        self.assertEqual(len(shared["client"].calls), 3)
        self.assertEqual(sorted(job.groups.values_list("attempts", flat=True)), [1, 2])
        self.assertEqual(job.error_message, "")

    def test_exhausted_attempts_mark_the_group_and_its_rows_failed(self):
        patch, shared = fake_gemini(raise_always=True)
        with patch:
            job = pipeline.run_audit(self.dt)
        self.assertEqual(job.status, AuditJob.Status.FAILED)
        self.assertEqual(len(shared["client"].calls), 4)  # مجموعتان × محاولتان
        self.assertEqual(set(job.groups.values_list("attempts", flat=True)), {2})
        self.assertEqual(set(job.groups.values_list("status", flat=True)), {"failed"})
        self.assertEqual(job.rows.filter(status="failed").count(), 10)
        self.assertEqual(job.rows.filter(status="skipped").count(), 2)
        self.assertIn("HTTP 503", job.error_message)
        self.assertEqual(job.finding_count, 3)  # ملاحظات القواعد وحدها
        self.assertEqual(job.stats["general"]["صفوف فشل فحصها"], 10)

    def test_partial_when_only_some_groups_fail(self):
        # ترتيب النداءات: الأولى ثم الثانية ثم المُعادة؛ الأولى تفشل في محاولتيها
        # والثانية تنجح من أول مرة
        patch, shared = fake_gemini(broken_calls={1, 3})
        with patch:
            job = pipeline.run_audit(self.dt)
        self.assertEqual(job.status, AuditJob.Status.PARTIAL)
        self.assertEqual(len(shared["client"].calls), 3)
        self.assertEqual(list(job.groups.values_list("attempts", flat=True)), [2, 1])
        self.assertEqual(sorted(job.groups.values_list("status", flat=True)), ["done", "failed"])
        self.assertEqual(job.rows.filter(status="failed").count(), 5)
        self.assertEqual(job.rows.filter(status="done").count(), 5)

    def test_no_api_key_fails_before_creating_a_job(self):
        with (
            override_settings(GEMINI_API_KEY=""),
            self.assertRaisesMessage(pipeline.AuditError, "GEMINI_API_KEY"),
        ):
            pipeline.run_audit(self.dt)
        self.assertFalse(AuditJob.objects.exists())

    def test_unexpected_error_is_recorded_on_the_job_and_raised(self):
        job = pipeline.create_job(self.dt)
        with (
            mock.patch.object(pipeline, "run_sync", side_effect=RuntimeError("انفجار")),
            self.assertRaisesMessage(RuntimeError, "انفجار"),
        ):
            pipeline.run_job(job)
        job.refresh_from_db()
        self.assertEqual(job.status, AuditJob.Status.FAILED)
        self.assertEqual(job.error_message, "RuntimeError: انفجار")

    def test_settings_come_from_django_with_overrides(self):
        loaded = load_settings(model="x", group_size=None)
        self.assertEqual((loaded.api_key, loaded.model, loaded.group_size), ("test-key", "x", 5))
        self.assertEqual(loaded.max_attempts, 2)

    def test_xlsx_export_has_the_four_sheets(self):
        patch, _shared = fake_gemini()
        with patch:
            job = pipeline.run_audit(self.dt)
        workbook = load_workbook(BytesIO(pipeline.export_xlsx(job)))
        self.assertEqual(workbook.sheetnames, ["الترجمات", "الأخطاء", "إحصائيات", "التصنيف"])
        translations = workbook["الترجمات"]
        headers = [c.value for c in translations[1]]
        self.assertEqual(
            headers[:5], ["#", "رقم الصف في الأصل", "النص العربي", "الترجمة", "معرّف الجملة"]
        )
        self.assertEqual(translations.max_row, 13)
        self.assertEqual(translations.cell(row=3, column=headers.index("الحكم") + 1).value, "مرفوض")
        self.assertEqual(workbook["الأخطاء"].max_row, 1 + job.finding_count)


class TaskTests(AuditFixture):
    def test_task_runs_a_pending_job_and_reports(self):
        job = pipeline.create_job(self.dt)
        patch, _shared = fake_gemini()
        with patch:
            result = tasks.audit_job_task.enqueue(job.pk)
        self.assertEqual(result.status, TaskResultStatus.SUCCESSFUL)
        self.assertEqual(result.return_value["status"], "done")
        self.assertEqual(result.return_value["findings"], 8)
        self.assertEqual(
            tasks.audit_job_task.enqueue(999999).return_value, {"job": 999999, "missing": True}
        )
        self.assertEqual(tasks.audit_job_task.enqueue(job.pk).return_value["skipped"], "done")

    def test_task_failure_marks_the_job(self):
        job = pipeline.create_job(self.dt)
        with mock.patch.object(pipeline, "run_sync", side_effect=RuntimeError("انفجار")):
            result = tasks.audit_job_task.enqueue(job.pk)
        self.assertEqual(result.status, TaskResultStatus.FAILED)
        job.refresh_from_db()
        self.assertEqual(job.status, AuditJob.Status.FAILED)


class ApiTests(AuditFixture):
    def run_url(self):
        return reverse("audit:run")

    def test_run_creates_and_executes_a_job(self):
        self.assertEqual(self.run_url(), "/api/audit/run/")
        patch, _shared = fake_gemini()
        with patch:
            response = self.client.post(
                self.run_url(), {"translation_id": self.dt.pk, "extra_rules": " قاعدة "}
            )
        self.assertEqual(response.status_code, 200)
        data = response.json()
        job = AuditJob.objects.get(pk=data["job_id"])
        self.assertEqual(
            (data["status"], data["finished"], data["translation_id"]), ("done", True, self.dt.pk)
        )
        self.assertEqual(job.extra_rules, "قاعدة")
        self.assertIsNone(job.created_by)
        self.assertEqual(self.client.get(self.run_url()).status_code, 405)

    def test_run_rejections(self):
        self.assertEqual(
            self.client.post(self.run_url(), {"translation_id": 999999}).status_code, 404
        )
        with override_settings(GEMINI_API_KEY=""):
            response = self.client.post(self.run_url(), {"translation_id": self.dt.pk})
        self.assertEqual(response.status_code, 503)
        pipeline.create_job(self.dt)  # بانتظار البدء
        response = self.client.post(self.run_url(), {"translation_id": self.dt.pk})
        self.assertEqual(response.status_code, 409)
        self.assertIn("جارٍ", response.json()["error"])

    def test_job_latest_and_xlsx_endpoints(self):
        latest = reverse("audit:latest", args=[self.dt.pk])
        self.assertEqual(
            self.client.get(latest).json(), {"translation_id": self.dt.pk, "job": None}
        )
        patch, _shared = fake_gemini()
        with patch:
            job = pipeline.run_audit(self.dt)
        data = self.client.get(reverse("audit:job", args=[job.pk])).json()["job"]
        self.assertEqual((data["job_id"], data["status"], data["finished"]), (job.pk, "done", True))
        self.assertEqual(data["counts"], {"rows": 12, "eligible": 10, "findings": 8})
        self.assertEqual(data["progress"]["groups"], {"done": 2})
        self.assertEqual(data["xlsx_url"], f"/api/audit/jobs/{job.pk}/xlsx/")
        self.assertEqual(len(data["rows"]), 12)
        row9 = next(r for r in data["rows"] if r["seq"] == 9)
        self.assertEqual(row9["phrase_id"], self.phrases[8].pk)
        self.assertEqual([f["type_code"] for f in row9["findings"]], ["SEM-OMI", "SEM-NEG"])
        self.assertEqual(row9["findings"][1]["explanation"], "نفي قُلب إثباتًا.")
        self.assertEqual(row9["status_label"], "تم")

        light = self.client.get(latest + "?rows=0").json()["job"]
        self.assertEqual(light["job_id"], job.pk)
        self.assertNotIn("rows", light)

        response = self.client.get(reverse("audit:xlsx", args=[job.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], XLSX_TYPE)
        self.assertIn(f"تدقيق-{job.pk}.xlsx", unquote(response["Content-Disposition"]))
        self.assertEqual(b"".join(response.streaming_content)[:2], b"PK")
        self.assertEqual(self.client.get(reverse("audit:job", args=[999999])).status_code, 404)


class CommandTests(AuditFixture):
    def test_audit_translation_runs_and_writes_excel(self):
        out, err = StringIO(), StringIO()
        output = Path(tempfile.mkdtemp(prefix="moneer-audit-")) / "نتائج.xlsx"
        patch, _shared = fake_gemini()
        with patch:
            call_command(
                "audit_translation", self.dt.pk, output=str(output), stdout=out, stderr=err
            )
        text = out.getvalue()
        self.assertIn("مكتمل", text)
        self.assertIn("الصفوف: 12 (مؤهلة للنموذج: 10) — الملاحظات: 8", text)
        self.assertIn("مرفوض: 6", text)
        self.assertIn("التوكن: إدخال 2000 / إخراج 400", text)
        self.assertTrue(output.exists())
        self.assertEqual(err.getvalue(), "")

    def test_command_errors(self):
        with self.assertRaisesMessage(CommandError, "لا توجد ترجمة"):
            call_command("audit_translation", 999999)
        with (
            override_settings(GEMINI_API_KEY=""),
            self.assertRaisesMessage(CommandError, "GEMINI_API_KEY"),
        ):
            call_command("audit_translation", self.dt.pk)


class AdminTests(AuditFixture):
    def test_models_are_registered_and_pages_render(self):
        for model in (AuditJob, AuditRow, AuditFinding):
            self.assertIn(model, admin.site._registry)
        patch, _shared = fake_gemini()
        with patch:
            job = pipeline.run_audit(self.dt)
        user = User.objects.create_superuser(
            email=f"{uuid.uuid4().hex[:8]}@example.com", password="x", full_name="م"
        )
        self.client.force_login(user)
        for name in ("auditjob", "auditrow", "auditfinding"):
            self.assertEqual(
                self.client.get(reverse(f"admin:audit_{name}_changelist")).status_code, 200
            )
        self.assertEqual(
            self.client.get(reverse("admin:audit_auditjob_change", args=[job.pk])).status_code, 200
        )
        row = job.rows.get(seq=9)
        self.assertEqual(
            self.client.get(reverse("admin:audit_auditrow_change", args=[row.pk])).status_code, 200
        )
